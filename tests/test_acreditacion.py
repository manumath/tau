"""Anexo de I+D+i para la acreditación (IAC-CINDA)."""
import io

import pytest
from openpyxl import load_workbook

from app.acreditacion import anexo as acreditacion
from app.investigacion import actividades, productos, proyectos, revistas
from app.comun import personas, usuarios, vinculos
from app.comun.db import get_db
from tests.conftest import PASS_ADMIN, PASS_TEMP, login


def _id(tabla, nombre):
    return get_db().q(f"SELECT id FROM {tabla} WHERE nombre = ?", [nombre])[0]["id"]


def _producto(uid, titulo, tipo="Artículo en revista", anio=2024, revista_id=None, autores=None, revision="arbitrado"):
    autores = autores or [f"i:{uid}"]
    d, errores = productos.preparar(dict(tipo_id=str(_id("tipo_producto", tipo)), titulo=titulo, anio=str(anio), mes="5",
                                         estado_publicacion="publicado", revision=revision,
                                         revista_id=str(revista_id or "")), autores, [])
    assert not errores, errores
    return productos.guardar(d, autores, uid, filiaciones=[1] * len(autores))


def _actividad(uid, titulo, tipo, fecha="2024-06-10", rol="expositor", ambito="internacional", pais="Chile"):
    return actividades.guardar(dict(tipo_id=_id("tipo_actividad", tipo), titulo=titulo, fecha_inicio=fecha, fecha_fin=None,
                                    modalidad="presencial", ambito=ambito, pais=pais), [(f"i:{uid}", rol)], uid)


@pytest.fixture
def datos(app):
    admin = usuarios.por_email("admin@test.cl")["id"]
    emmy = usuarios.por_email("inv@test.cl")["id"]
    usuarios.cambiar_password(emmy, PASS_TEMP, temporal=False)
    vinculos.agregar(emmy, "contratado", "2019-03-01", "2022-12-31", 40)
    vinculos.agregar(emmy, "titular", "2023-01-01", None, 40)
    alan = usuarios.crear("alan@test.cl", "Alan", "Turing", PASS_TEMP)["id"]
    vinculos.agregar(alan, "invitado", "2024-01-01", "2024-12-31", 20)
    nacional = revistas.crear({"titulo": "Revista Boliviana", "pais": "Bolivia", "arbitrada": 1, "acceso_abierto": 1},
                              admin, verificada=True)
    internacional = revistas.crear({"titulo": "Annals", "pais": "Estados Unidos"}, admin, verificada=True)
    revistas.agregar_indexacion(internacional, {"indice_id": _id("indice", "Scopus"), "desde": "2000", "cuartil": "Q1"}, admin)
    pendiente = revistas.crear({"titulo": "Revista Nueva", "pais": "Perú"}, admin)
    weyl = personas.crear_externo("Hermann", "Weyl", "IAS", "Estados Unidos", emmy)["id"]
    _producto(emmy, "Nacional", revista_id=nacional)
    _producto(emmy, "Internacional con coautor", revista_id=internacional, autores=[f"i:{emmy}", f"e:{weyl}"])
    _producto(emmy, "Declarado", revista_id=pendiente)
    _producto(emmy, "Antes de la ventana", anio=2019, revista_id=nacional)
    _producto(emmy, "Un libro", tipo="Libro")
    _producto(emmy, "Una patente", tipo="Patente", revision="sin_arbitraje")
    _producto(alan, "Software de Alan", tipo="Software", revision="sin_arbitraje")
    _actividad(emmy, "Congreso con charla", "Congreso")
    _actividad(emmy, "Congreso de oyente", "Congreso", rol="asistente")
    _actividad(emmy, "Estancia en Santiago", "Estancia o pasantía de investigación")
    _actividad(emmy, "Asesoría al ministerio", "Transferencia de conocimiento", ambito="nacional", pais="Bolivia")
    _actividad(emmy, "Red andina", "Participación en red de investigación")
    _actividad(emmy, "Comité editorial", "Comité científico o editorial")
    proyectos.crear(dict(nombre="Proyecto IDH", estado="vigente", fecha_inicio="2023-10-01", fecha_fin="2024-09-30",
                         fuente_financiamiento="Recursos IDH – UMSA"), emmy, emmy)
    return {"emmy": emmy, "alan": alan}


def test_anexo(datos):
    a = acreditacion.anexo(2022, 2026)
    assert {p["titulo"] for p in a["articulos"]} == {"Nacional", "Internacional con coautor"}  # sin la declarada
    assert [p["titulo"] for p in a["memoria"]["por_confirmar"]] == ["Declarado"]
    assert [p["titulo"] for p in a["libros"]] == ["Un libro"]
    assert {p["titulo"] for p in a["tecnicos"]} == {"Una patente", "Software de Alan"}
    assert [x["titulo"] for x in a["presentaciones"]] == ["Congreso con charla"]  # solo con exposición
    assert [x["titulo"] for x in a["estancias"]] == ["Estancia en Santiago"] and len(a["movilidad_internacional"]) == 1
    assert [x["titulo"] for x in a["transferencia"]] == ["Asesoría al ministerio"]
    assert [x["titulo"] for x in a["redes"]] == ["Red andina"] and [x["titulo"] for x in a["comites"]] == ["Comité editorial"]
    assert [p["nombre"] for p in a["proyectos"]] == ["Proyecto IDH"] and a["proyectos"][0]["n_internos"] == 1
    assert (a["coautoria"]["internacionales"], a["coautoria"]["con_pais"], a["coautoria"]["por_pais"]) == (
        1, 1, {"Estados Unidos": 1})
    anio = {x["anio"]: x for x in a["anios"]}
    assert (anio[2024]["indexados"], anio[2024]["arbitrados"], anio[2024]["nacionales"], anio[2024]["internacionales"]) \
        == (1, 1, 1, 1)
    assert anio[2024]["tecnicos"] == 2 and anio[2024]["presentaciones"] == 1 and anio[2025]["indexados"] == 0
    # Sin exigir revista verificada, el artículo declarado cuenta.
    assert "Declarado" in {p["titulo"] for p in acreditacion.anexo(2022, 2026, exigir_verificada=False)["articulos"]}


def test_cuerpo_academico(datos):
    c = {x["anio"]: x for x in acreditacion.anexo(2022, 2026)["cuerpo"]}
    assert c[2022]["por_categoria"] == {"Contratado(a)": 1} and c[2022]["jce"] == 1.0
    assert c[2024]["por_categoria"] == {"Titular": 1, "Invitado(a)": 1} and (c[2024]["horas"], c[2024]["jce"]) == (240, 1.5)  # jornada × 4 (h/mes)
    assert c[2025]["total"] == 1  # el invitado terminó en 2024


def test_rutas(client, datos):
    login(client, "inv@test.cl", PASS_TEMP)
    assert client.get("/informes/acreditacion").status_code == 403
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "Anexo de I+D+i para la acreditación" in client.get("/informes").text
    r = client.get("/informes/acreditacion?desde=2022&hasta=2026")
    for texto in ("Resumen por año", "3.6.3 a", "Publicaciones en revistas indexadas", "Congreso con charla",
                  "Jornada completa equivalente", "Aún no registrado en el portal", "Datos y método", "por confirmar"):
        assert texto in r.text, texto
    r = client.get("/informes/acreditacion/xlsx?desde=2022&hasta=2026")
    assert r.headers["content-disposition"] == 'attachment; filename="anexo-idi-2022-2026.xlsx"'
    libro = load_workbook(io.BytesIO(r.content))
    assert "3.6.3 a Publicaciones" in libro.sheetnames and "Método" in libro.sheetnames
    resumen = libro["Resumen por año"]
    assert [c.value for c in resumen[1]] == ["Indicador", "2022", "2023", "2024", "2025", "2026"]
    # Una ventana de más de 10 años se acota.
    assert "Anexo de I+D+i para la acreditación, 2016–2026" in client.get("/informes/acreditacion?desde=1990&hasta=2026").text


def test_producto_del_anio_en_curso_sin_mes_cuenta(datos):
    """La ventana usa años completos: un artículo de este año sin mes no queda «por confirmar»."""
    from datetime import date
    anio = date.today().year
    admin = usuarios.por_email("admin@test.cl")["id"]
    rid = revistas.crear({"titulo": "Revista X", "pais": "Bolivia", "arbitrada": 1}, admin, verificada=True)
    d, _ = productos.preparar(dict(tipo_id=str(_id("tipo_producto", "Artículo en revista")), titulo="De este año",
                                   anio=str(anio), mes="", estado_publicacion="publicado", revision="arbitrado",
                                   revista_id=str(rid)), [f"i:{datos['emmy']}"], [])
    productos.guardar(d, [f"i:{datos['emmy']}"], datos["emmy"], filiaciones=[1])
    a = acreditacion.anexo(anio - 4, anio)
    assert "De este año" in {p["titulo"] for p in a["articulos"]}
