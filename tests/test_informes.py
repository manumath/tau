import io

import pytest
from openpyxl import load_workbook

from app.acreditacion import informes
from app.investigacion import perfiles, productos, proyectos
from app import personas, usuarios
from app.db import get_db
from tests.conftest import PASS_ADMIN, PASS_TEMP, login

CLAVE = "Clave-colega-1"


def _tipo(nombre):
    return get_db().q("SELECT id FROM tipo_producto WHERE nombre = ?", [nombre])[0]["id"]


def _indice(nombre):
    return get_db().q("SELECT id FROM indice WHERE nombre = ?", [nombre])[0]["id"]


def _producto(creador, autores, titulo, anio=2026, tipo="Artículo en revista", revision="arbitrado",
              estado="publicado", indices=(), filiacion=1, **extra):
    datos = dict(tipo_id=str(_tipo(tipo)), titulo=titulo, anio=str(anio), estado_publicacion=estado,
                 revision=revision, medio="Revista", volumen=None, numero=None, paginas=None, doi=None,
                 arxiv_id=None, url=None, resumen=None, actividad_id=None, **extra)
    d, errores = productos.preparar(datos, autores, [str(_indice(i)) for i in indices])
    assert not errores, errores
    # Por omisión, los autores del instituto firmaron con su filiación (si no, no cuenta).
    return productos.guardar(d, autores, creador, filiaciones=[filiacion] * len(autores))


@pytest.fixture
def datos(app):
    emmy = usuarios.por_email("inv@test.cl")["id"]
    usuarios.cambiar_password(emmy, PASS_TEMP, temporal=False)
    alan = usuarios.crear("alan@test.cl", "Alan", "Turing", CLAVE, debe_cambiar_pass=False,
                          categoria="contratado")["id"]
    # Vínculo con el instituto (sin fecha de ingreso: «desde siempre»); sin él, nada cuenta.
    perfiles.cambiar_categoria(emmy, "titular")
    weyl = personas.crear_externo("Hermann", "Weyl", "IAS", None, emmy)["id"]
    ids = {
        # Coautoría de dos investigadores: en la memoria debe contarse una sola vez.
        "indexado": _producto(emmy, [f"i:{emmy}", f"i:{alan}", f"e:{weyl}"], "Anillos noetherianos",
                              revision="indexado", indices=["Scopus", "MathSciNet"]),
        "arbitrado": _producto(alan, [f"i:{alan}"], "Máquinas", estado="aceptado"),
        "enviado": _producto(emmy, [f"i:{emmy}"], "En revisión", estado="enviado"),
        "reporte": _producto(emmy, [f"i:{emmy}"], "Notas", tipo="Reporte de investigación"),
        "viejo": _producto(emmy, [f"i:{emmy}"], "Antiguo", anio=2019),
        "otro_anio": _producto(emmy, [f"i:{emmy}"], "Del año pasado", anio=2025, revision="indexado",
                               indices=["Scopus"]),
        # Solo externos: no es producción del instituto.
        "ajeno": _producto(emmy, [f"e:{weyl}"], "Solo de Weyl"),
    }
    return {"emmy": emmy, "alan": alan, "ids": ids}


def test_categorias_del_cv():
    base = {"estado_publicacion": "publicado", "tipo": "Artículo en revista"}
    assert informes.categoria({**base, "revision": "indexado"}) == "indexados"
    assert informes.categoria({**base, "revision": "arbitrado"}) == "arbitrados"
    assert informes.categoria({**base, "revision": "indexado", "estado_publicacion": "enviado"}) == "en_curso"
    assert informes.categoria({**base, "tipo": "Libro", "revision": "arbitrado"}) == "libros"
    assert informes.categoria({**base, "tipo": "Preprint", "revision": "sin_arbitraje"}) == "preprints"


def test_periodo_valido():
    assert informes.periodo_valido("2026", "2020") == (2020, 2026)
    assert informes.periodo_valido("x", "")[1] - informes.periodo_valido("x", "")[0] == 4


def test_cv(datos):
    cv = informes.cv(datos["emmy"], 2020, 2026)
    titulos = {k: [p["titulo"] for p in v] for k, v in cv["grupos"].items() if v}
    assert titulos == {"indexados": ["Anillos noetherianos", "Del año pasado"],
                       "en_curso": ["En revisión"], "reportes": ["Notas"]}
    # «Solo de Weyl» lo registró Emmy pero no es autora: no va en su CV. 2019 queda fuera del periodo.
    assert cv["n_productos"] == 4 and cv["n_informe"] == 2
    assert informes.cv(datos["emmy"], 2020, 2026, solo_informe=True)["n_productos"] == 2


def test_pagina_cv(client, datos):
    login(client, "inv@test.cl", PASS_TEMP)
    r = client.get("/informes/cv?desde=2026&hasta=2026")
    assert "Artículos en revistas indexadas (1)" in r.text and "Anillos noetherianos" in r.text
    assert "<strong>Noether, E.</strong>" in r.text  # la persona del CV, destacada
    assert "Scopus, MathSciNet" in r.text and "window.print()" in r.text
    assert "Del año pasado" not in r.text


def test_cv_de_otro_solo_para_admin(client, datos):
    login(client, "alan@test.cl", CLAVE)
    r = client.get(f"/informes/cv?uid={datos['emmy']}")
    assert "Alan Turing" in r.text and "Emmy Noether" not in r.text  # ignora uid ajeno
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "Emmy Noether" in client.get(f"/informes/cv?uid={datos['emmy']}").text


def test_memoria(datos):
    m = informes.memoria(2026)
    assert [p["titulo"] for p in m["productos"]] == ["Anillos noetherianos", "Máquinas"]
    assert (m["n_indexados"], m["n_arbitrados"]) == (1, 1)
    assert m["por_indice"] == {"Scopus": 1, "MathSciNet": 1}
    por_inv = {x["apellidos"]: x["productos"] for x in m["por_investigador"]}
    assert por_inv == {"Noether": 1, "Turing": 2}  # el artículo común cuenta para ambos…
    assert len(m["productos"]) == 2                 # …pero una sola vez en el total


def test_memoria_solo_admin_y_descargas(client, datos):
    login(client, "inv@test.cl", PASS_TEMP)
    assert client.get("/informes/memoria").status_code == 403
    assert client.get("/informes/memoria/xlsx").status_code == 403
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "Memoria anual 2026" in client.get("/informes/memoria?anio=2026").text

    r = client.get("/informes/memoria/xlsx?anio=2026")
    assert r.headers["content-disposition"] == 'attachment; filename="memoria-2026.xlsx"'
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Productos", "Actividades", "Proyectos", "Por categoría", "Por investigador"]
    hoja = libro["Productos"]
    assert [c.value for c in hoja[1]][:4] == ["Año", "Mes", "Tipo", "Título"] and hoja.max_row == 3

    r = client.get("/informes/memoria/csv/productos?anio=2026")
    assert r.content.startswith(b"\xef\xbb\xbf")  # BOM para Excel
    texto = r.content.decode("utf-8-sig")
    assert "Anillos noetherianos" in texto and "Noether, Emmy; Turing, Alan" in texto
    assert client.get("/informes/memoria/csv/otra").status_code == 404


def test_inyeccion_de_formulas_neutralizada(client, datos):
    _producto(datos["emmy"], [f"i:{datos['emmy']}"], "=HYPERLINK(\"http://malo\")")
    assert informes.celda_segura("=1+1") == "'=1+1" and informes.celda_segura("-x") == "'-x"
    assert informes.celda_segura(5) == 5 and informes.celda_segura("Normal") == "Normal"
    login(client, "admin@test.cl", PASS_ADMIN)
    libro = load_workbook(io.BytesIO(client.get("/informes/memoria/xlsx?anio=2026").content))
    valores = [c.value for fila in libro["Productos"].iter_rows() for c in fila]
    assert "'=HYPERLINK(\"http://malo\")" in valores
    assert all(not (isinstance(v, str) and v.startswith("=")) for v in valores)


def test_informe_de_proyecto(client, datos):
    pid = proyectos.crear(dict(nombre="Álgebra", estado="vigente", fecha_inicio="2025-03-01",
                               objetivo_general="Estudiar anillos."), datos["emmy"], datos["emmy"])
    proyectos.agregar_objetivo(pid, "Primer objetivo")
    for clave in ("indexado", "reporte"):
        proyectos.sincronizar("producto", datos["ids"][clave], [pid], {"id": datos["emmy"]})
    # Propuesto por otra persona y sin aprobar: no aparece en el informe.
    proyectos.agregar_participante(pid, f"i:{datos['alan']}", "co_investigador", True)
    proyectos.sincronizar("producto", datos["ids"]["arbitrado"], [pid], {"id": datos["alan"]})
    inf = informes.informe_proyecto(pid, 2025, 2026)
    assert [p["titulo"] for p in inf["productos_informe"]] == ["Anillos noetherianos"]
    assert [p["titulo"] for p in inf["productos_otros"]] == ["Notas"]

    login(client, "inv@test.cl", PASS_TEMP)
    r = client.get(f"/proyectos/{pid}/informe")
    assert "Informe de avance · 2025–" in r.text and "Primer objetivo" in r.text and "Máquinas" not in r.text
    r = client.get(f"/proyectos/{pid}/informe/xlsx")
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Productos", "Actividades"]
    client.post("/logout")
    usuarios.crear("sofia@test.cl", "Sofía", "K", CLAVE, debe_cambiar_pass=False)
    login(client, "sofia@test.cl", CLAVE)
    assert client.get(f"/proyectos/{pid}/informe/xlsx").status_code == 404


def test_indice_de_informes(client, datos):
    login(client, "inv@test.cl", PASS_TEMP)
    r = client.get("/informes")
    assert "Mi CV" in r.text and "Memoria del instituto" not in r.text
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "Memoria del instituto" in client.get("/informes").text


def test_memoria_por_categoria(datos):
    from app.investigacion import actividades, perfiles
    perfiles.cambiar_categoria(datos["emmy"], "titular")
    perfiles.cambiar_categoria(datos["alan"], "contratado")
    tipo = get_db().q("SELECT id FROM tipo_actividad ORDER BY orden LIMIT 1")[0]["id"]
    actividades.guardar(dict(tipo_id=tipo, titulo="Seminario", fecha_inicio="2026-05-04", fecha_fin=None,
                             modalidad="presencial", ambito="nacional"),
                        [(f"i:{datos['emmy']}", "expositor")], datos["emmy"])
    m = informes.memoria(2026)
    assert m["por_categoria"]["Titular"] == {"investigadores": 1, "productos": 1, "actividades": 1}
    assert m["por_categoria"]["Contratado(a)"] == {"investigadores": 1, "productos": 2, "actividades": 0}
    assert {x["apellidos"]: x["categoria"] for x in m["por_investigador"]} == \
        {"Noether": "titular", "Turing": "contratado"}


def test_memoria_externos_solo_con_proyecto_institucional(datos):
    """Un investigador de categoría «Externo» figura solo con lo ligado a un proyecto institucional."""
    from app.investigacion import perfiles
    externo = usuarios.crear("ext@test.cl", "Hermann", "Minkowski", CLAVE, categoria="externo")["id"]
    solo = _producto(externo, [f"i:{externo}"], "Solo del externo")
    ligado = _producto(externo, [f"i:{externo}"], "Del externo en proyecto")
    perfiles.cambiar_categoria(datos["emmy"], "titular")

    m = informes.memoria(2026)
    titulos = {p["titulo"] for p in m["productos"]}
    assert "Solo del externo" not in titulos and "Del externo en proyecto" not in titulos
    assert "Minkowski" not in {x["apellidos"] for x in m["por_investigador"]}

    # Proyecto individual: no alcanza.
    indiv = proyectos.crear(dict(nombre="Individual", estado="vigente", fecha_inicio="2026-01-01"),
                            externo, externo)
    proyectos.sincronizar("producto", ligado, [indiv], {"id": externo})
    assert "Del externo en proyecto" not in {p["titulo"] for p in informes.memoria(2026)["productos"]}

    # Proyecto institucional vigente en el año: figura (y el externo aparece en el resumen).
    inst = proyectos.crear(dict(nombre="Institucional", estado="vigente", fecha_inicio="2025-10-01",
                                fecha_fin="2026-09-30"), datos["emmy"], datos["emmy"], tipo="institucional")
    proyectos.agregar_participante(inst, f"i:{externo}", "co_investigador", False)
    proyectos.sincronizar("producto", ligado, [inst], {"id": externo})
    assert proyectos.vinculo("producto", ligado, inst)["estado"] == "pendiente"
    assert "Del externo en proyecto" not in {p["titulo"] for p in informes.memoria(2026)["productos"]}
    proyectos.aprobar("producto", ligado, inst, datos["emmy"])
    m = informes.memoria(2026)
    titulos = {p["titulo"] for p in m["productos"]}
    assert "Del externo en proyecto" in titulos and "Solo del externo" not in titulos
    assert {x["apellidos"]: x["productos"] for x in m["por_investigador"]}["Minkowski"] == 1
    assert m["por_categoria"]["Externo(a)"] == {"investigadores": 1, "productos": 1, "actividades": 0}
    # El mismo proyecto no alcanza a otro año.
    assert "Del externo en proyecto" not in {p["titulo"] for p in informes.memoria(2027)["productos"]}
    assert [p["tipo"] for p in m["proyectos"]][0] == "institucional"


def test_memoria_producto_mixto_cuenta_por_el_autor_no_externo(datos):
    externo = usuarios.crear("ext@test.cl", "Hermann", "Minkowski", CLAVE, categoria="externo")["id"]
    _producto(datos["emmy"], [f"i:{datos['emmy']}", f"i:{externo}"], "Mixto")
    m = informes.memoria(2026)
    assert "Mixto" in {p["titulo"] for p in m["productos"]}
    # Cuenta para Emmy, no para el externo (no está ligado a un proyecto institucional).
    assert "Minkowski" not in {x["apellidos"] for x in m["por_investigador"]}
