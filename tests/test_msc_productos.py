"""Clasificación MSC2020 de los productos (PLAN.md §45): un código principal y hasta 4 secundarios,
sugerida desde zbMATH Open o el perfil, con filtro, memoria por área y calidad de datos."""
import io
import json

import httpx2
import pytest
from openpyxl import load_workbook

from app.acreditacion import calidad, informes
from app.investigacion import importacion, msc, perfiles, productos
from app.comun.db import get_db
from tests.conftest import HX, PASS_ADMIN, login
from tests.test_importacion import CROSSREF
from tests.test_informes import datos  # noqa: F401  (fixture: productos de 2026 que cuentan)
from tests.test_productos import _form, inv  # noqa: F401


def _codigos(pid):
    return [m["codigo"] for m in msc.de_producto(pid)]


def _nuevo(cliente, **extra):
    r = cliente.post("/productos", headers=HX, data=_form(cliente.uid, **extra))
    assert r.status_code == 200 and "Producto registrado" in r.text, r.text
    return get_db().q("SELECT max(id) AS id FROM producto")[0]["id"]


# --- Reglas ------------------------------------------------------------------------------------

def test_validar(app):
    assert msc.validar(["11M06", " 11-XX ", ""]) == (["11M06", "11-XX"], [])
    assert "repetidos" in msc.validar(["11M06", "11M06"])[1][0]
    assert "a lo sumo 5" in msc.validar(["11-XX", "13-XX", "14-XX", "15-XX", "16-XX", "17-XX"])[1][0]
    assert "no existen" in msc.validar(["99Z99"])[1][0]


def test_guardar_ordenar_y_quitar(inv):
    pid = _nuevo(inv, msc=["13E05", "11-XX"])
    assert _codigos(pid) == ["13E05", "11-XX"]
    p = productos.obtener(pid)
    assert p["msc"][0]["principal"] and not p["msc"][1]["principal"]
    # Al editar, el orden del formulario define el principal; sin códigos, queda sin clasificar.
    assert inv.put(f"/productos/{pid}", headers=HX, data=_form(inv.uid, msc=["11-XX", "13E05"])).status_code == 200
    assert _codigos(pid) == ["11-XX", "13E05"]
    assert "Cambios guardados" in inv.put(f"/productos/{pid}", headers=HX, data=_form(inv.uid)).text
    assert _codigos(pid) == []


def test_errores_conservan_lo_elegido(inv):
    r = inv.post("/productos", headers=HX, data=_form(inv.uid, msc=["13E05", "13E05"]))
    assert "repetidos" in r.text and 'value="13E05"' in r.text
    assert not get_db().q("SELECT 1 FROM producto")


def test_otras_vias_no_tocan_la_clasificacion(inv):
    """BibTeX, OpenAlex y el seed no traen `msc`: guardar sin ella no borra la que había."""
    pid = _nuevo(inv, msc=["13E05"])
    sin_msc = {k: v for k, v in _form(inv.uid, titulo="Otro título").items() if k != "autor"}
    d, errores = productos.preparar(sin_msc, [f"i:{inv.uid}"], [])
    assert not errores and "_msc" not in d
    productos.guardar(d, [f"i:{inv.uid}"], inv.uid, pid)
    assert _codigos(pid) == ["13E05"]


def test_eliminar_el_producto_borra_su_clasificacion(inv):
    pid = _nuevo(inv, msc=["13E05"])
    productos.eliminar(pid)
    assert not get_db().q("SELECT 1 FROM producto_msc")


# --- Formulario ---------------------------------------------------------------------------------

def test_buscador_y_fila(inv):
    r = inv.get("/productos/msc/buscar", params={"msc_q": "13E", "msc": ["13E05"]}, headers=HX)
    assert "13E10" in r.text and ">13E05<" not in r.text and "data-tras-pedido-limpiar" in r.text
    llenos = ["11-XX", "13-XX", "14-XX", "15-XX", "16-XX"]
    assert "quite uno" in inv.get("/productos/msc/buscar", params={"msc_q": "zeta", "msc": llenos}, headers=HX).text
    fila = inv.get("/productos/msc/fila", params={"codigo": "13E05"}, headers=HX).text
    assert 'name="msc"' in fila and 'value="13E05"' in fila and "group-first:inline-flex" in fila
    assert inv.get("/productos/msc/fila", params={"codigo": "99Z99"}, headers=HX).text == ""


def test_nuevo_sugiere_el_area_principal_del_perfil(inv):
    assert "Sugerida según" not in inv.get("/productos/nuevo", headers=HX).text
    perfiles.agregar_area(inv.uid, "13-XX")
    perfiles.agregar_area(inv.uid, "11M06")
    r = inv.get("/productos/nuevo", headers=HX).text
    assert 'value="13-XX"' in r and 'value="11M06"' not in r and "Sugerida según su perfil" in r


class Red:
    """Crossref responde con CROSSREF; zbMATH, con `zbmath` (o 404 si es None)."""
    def __init__(self, zbmath=None):
        self.zbmath, self.llamadas = zbmath, []

    def __call__(self, url, params=None):
        self.llamadas.append((url, params))
        if "zbmath" in url:
            if self.zbmath is None:
                return httpx2.Response(404, content=b'{"result": null}')
            return httpx2.Response(200, content=json.dumps({"result": [{"msc": self.zbmath}]}).encode())
        return httpx2.Response(200, content=json.dumps(CROSSREF).encode())


def test_importar_sugiere_la_clasificacion_de_zbmath(inv, monkeypatch):
    red = Red([{"code": "68Q15", "scheme": "msc2020"}, {"code": "99Z99", "scheme": "msc2020"},
               {"code": "03D15", "scheme": "msc2020"}])
    monkeypatch.setattr(importacion, "_get", red)
    r = inv.post("/productos/importar", headers=HX, data={"identificador": "10.1090/x"}).text
    assert 'value="68Q15"' in r and 'value="03D15"' in r and "99Z99" not in r  # solo códigos del catálogo
    assert r.index('value="68Q15"') < r.index('value="03D15"') and "Sugerida según zbMATH Open" in r
    assert red.llamadas[-1][1]["DOI"] == "10.1090/x"


def test_importar_sin_zbmath_usa_el_perfil_de_los_autores(inv, monkeypatch):
    monkeypatch.setattr(importacion, "_get", Red(None))
    perfiles.agregar_area(inv.uid, "13-XX")  # Emmy Noether figura entre los autores de CROSSREF
    r = inv.post("/productos/importar", headers=HX, data={"identificador": "10.1090/x"}).text
    assert 'value="13-XX"' in r and "Sugerida según el perfil de los autores" in r


def test_zbmath_caido_no_estorba(app, monkeypatch):
    def caido(url, params=None):
        raise importacion.ErrorImportacion("sin conexión")
    monkeypatch.setattr(importacion, "_get", caido)
    assert importacion.msc_de_zbmath("10.1090/x") == []
    monkeypatch.setattr(importacion, "_get", lambda url, params=None: httpx2.Response(200, content=b"<html>"))
    assert importacion.msc_de_zbmath("10.1090/x") == []
    assert importacion.msc_de_zbmath(None) == []


# --- Ficha, lista y público -------------------------------------------------------------------------

def test_ficha_lista_y_filtro(inv):
    algebra = _nuevo(inv, titulo="Anillos", msc=["13E05", "11-XX"])
    _nuevo(inv, titulo="Funciones zeta", msc=["11M06"])
    _nuevo(inv, titulo="Sin clasificar")
    ficha = inv.get(f"/productos/{algebra}").text
    assert "Clasificación MSC2020" in ficha and "Commutative Noetherian rings and modules" in ficha
    pagina = inv.get("/productos").text
    assert 'name="area"' in pagina and 'value="11"' in pagina and 'value="13"' in pagina and "MSC 13E05" in pagina
    # La sección cuenta en el código principal o en uno secundario.
    once = inv.get("/productos/lista", params={"area": "11"}, headers=HX)
    assert "Anillos" in once.text and "Funciones zeta" in once.text and "Sin clasificar" not in once.text
    assert once.headers["hx-replace-url"] == "/productos?area=11"
    trece = inv.get("/productos/lista", params={"area": "13"}, headers=HX).text
    assert "Anillos" in trece and "Funciones zeta" not in trece
    assert "Funciones zeta" not in inv.get("/productos/bibtex", params={"area": "13"}).text


def test_api_publica(inv):
    get_db().execute("UPDATE investigador SET perfil_publico = 1 WHERE usuario_id = ?", [inv.uid])
    _nuevo(inv, titulo="Anillos", msc=["13E05", "11-XX"])
    api = inv.get("/api/v1/productos").json()
    assert api[0]["msc"] == [{"codigo": "13E05", "texto": "Commutative Noetherian rings and modules",
                                "principal": True},
                               {"codigo": "11-XX", "texto": "Number theory", "principal": False}]
    assert "MSC2020: 13E05, 11-XX" in inv.get(f"/publico/investigadores/{inv.uid}").text


# --- Memoria y calidad de datos -----------------------------------------------------------------------

def test_memoria_por_area(client, datos):  # noqa: F811
    ids = datos["ids"]
    msc.guardar_de_producto(ids["indexado"], ["13E05", "68Q70"])
    m = informes.memoria(2026)
    assert m["por_area"] == [{"seccion": "13", "indexados": 1, "arbitrados": 0,
                              "texto": "Commutative algebra", "total": 1}]
    assert m["sin_area"] == 1  # «Máquinas», arbitrado, sin clasificar
    sin_msc = {c["clave"]: c for c in calidad.revisar()}["sin_msc"]["items"]  # la ventana abarca varios años
    assert sorted(i["texto"] for i in sin_msc) == ["Del año pasado (2025)", "Máquinas (2026)"]
    login(client, "admin@test.cl", PASS_ADMIN)
    r = client.get("/informes/memoria", params={"anio": 2026})
    assert "Por área (MSC2020)" in r.text and "Productos por sección MSC2020" in r.text and "Sin clasificar" in r.text
    libro = load_workbook(io.BytesIO(client.get("/informes/memoria/xlsx", params={"anio": 2026}).content))
    hoja = libro["Productos"]
    columnas = [c.value for c in hoja[1]]
    valores = {fila[columnas.index("Título")]: fila[columnas.index("MSC2020 (principal primero)")]
               for fila in hoja.iter_rows(min_row=2, values_only=True)}
    assert valores["Anillos noetherianos"] == "13E05, 68Q70"


@pytest.mark.parametrize("codigo,seccion", [("11M06", "11"), ("11-XX", "11"), ("03D80", "03")])
def test_seccion_del_principal(codigo, seccion, app):
    assert msc.seccion(codigo)["codigo"] == f"{seccion}-XX"
