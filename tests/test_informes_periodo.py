"""Ajustes, etapa D: mes de publicación, memoria por periodo o rango, plan de actividades y
exportación de convocatorias."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import actividades, bibtex, importacion, informes
from app import institucionales as inst
from app import perfiles, productos, proyectos, usuarios
from app.db import get_db
from tests.conftest import HX, PASS_ADMIN, PASS_TEMP, login
from tests.test_importacion import red  # noqa: F401  (fixture con respuestas simuladas)
from tests.test_propuestas import COMPLETA


def _producto(uid, titulo, anio, mes=None):
    tipo = get_db().q("SELECT id FROM tipo_producto WHERE nombre = 'Artículo en revista'")[0]["id"]
    d, errores = productos.preparar(dict(tipo_id=str(tipo), titulo=titulo, anio=str(anio), mes=str(mes or ""),
                                         estado_publicacion="publicado", revision="arbitrado"), [f"i:{uid}"], [])
    assert not errores, errores
    return productos.guardar(d, [f"i:{uid}"], uid)


def _actividad(uid, titulo, fecha):
    tipo = get_db().q("SELECT id FROM tipo_actividad ORDER BY orden LIMIT 1")[0]["id"]
    return actividades.guardar(dict(tipo_id=tipo, titulo=titulo, fecha_inicio=fecha, fecha_fin=None,
                                    modalidad="presencial", ambito="nacional"), [(f"i:{uid}", "expositor")], uid)


@pytest.fixture
def emmy(app):
    uid = usuarios.por_email("inv@test.cl")["id"]
    usuarios.cambiar_password(uid, PASS_TEMP, temporal=False)
    perfiles.cambiar_categoria(uid, "titular")
    return uid


# --- Mes de publicación --------------------------------------------------------------------------------

def test_mes_de_publicacion(app, emmy):
    tipo = get_db().q("SELECT id FROM tipo_producto WHERE nombre = 'Artículo en revista'")[0]["id"]
    base = dict(tipo_id=str(tipo), titulo="X", anio="2026", estado_publicacion="publicado", revision="arbitrado")
    assert productos.preparar({**base, "mes": "13"}, [f"i:{emmy}"], [])[1] == ["Mes de publicación no válido."]
    d, _ = productos.preparar({**base, "mes": ""}, [f"i:{emmy}"], [])
    assert d["mes"] is None
    pid = _producto(emmy, "Con mes", 2026, 11)
    assert productos.obtener(pid)["mes"] == 11


def test_mes_en_el_formulario_y_la_ficha(client, emmy):
    login(client, "inv@test.cl", PASS_TEMP)
    assert 'name="mes"' in client.get("/productos/nuevo", headers=HX).text
    pid = _producto(emmy, "Con mes", 2026, 3)
    assert "marzo 2026" in client.get(f"/productos/{pid}").text


def test_mes_al_importar(app, red):  # noqa: F811
    assert importacion.importar("10.1090/x")["datos"]["mes"] == 5            # Crossref: [[1965, 5]]
    assert importacion.importar("arXiv:2401.01234")["datos"]["mes"] == 1     # arXiv: 2024-01-03


@pytest.mark.parametrize("campo,mes", [("month = mar", "3"), ("month = {November}", "11"), ("month = 7", "7"),
                                       ("date = {2024-02-10}", "2"), ("", "")])
def test_mes_en_bibtex(app, campo, mes):
    entradas, _ = bibtex.leer(f"@article{{x, author={{Noether, E.}}, title={{T}}, journal={{J}}, year={{2024}}, {campo}}}")
    assert entradas[0]["datos"]["mes"] == mes


def test_bibtex_exporta_el_mes(app, emmy):
    _producto(emmy, "Con mes", 2026, 4)
    assert "month = {4}" in bibtex.exportar(productos.de_investigador(emmy))


# --- Memoria por periodo y por rango ------------------------------------------------------------------------

@pytest.mark.parametrize("anio,mes,esperado", [(2026, 11, "dentro"), (2026, 9, "fuera"), (2027, 9, "dentro"),
                                               (2027, 10, "fuera"), (2026, None, "dudoso"), (2025, None, "fuera")])
def test_ubicar_producto_en_periodo(anio, mes, esperado):
    assert informes.ubicar_producto({"anio": anio, "mes": mes}, "2026-10-01", "2027-09-30") == esperado


def test_rangos():
    assert informes.rango(2026)["titulo"] == "Memoria anual 2026"
    r = informes.rango(periodo=2026)
    assert (r["desde"], r["hasta"], r["params"]) == ("2026-10-01", "2027-09-30", {"periodo": 2026})
    assert informes.rango(desde="2026-03-01", hasta="2026-06-30")["titulo"] == "Memoria del 1 mar 2026 al 30 jun 2026"


def test_memoria_por_periodo(app, emmy):
    _producto(emmy, "Noviembre 2026", 2026, 11)
    _producto(emmy, "Marzo 2026", 2026, 3)
    _producto(emmy, "Sin mes 2026", 2026)
    _producto(emmy, "Mayo 2027", 2027, 5)
    _actividad(emmy, "Seminario de octubre", "2026-10-15")
    _actividad(emmy, "Congreso de agosto", "2026-08-20")

    m = informes.memoria(periodo=2026)
    assert {p["titulo"] for p in m["productos"]} == {"Noviembre 2026", "Mayo 2027"}
    assert [p["titulo"] for p in m["por_confirmar"]] == ["Sin mes 2026"]
    assert [a["titulo"] for a in m["actividades"]] == ["Seminario de octubre"]
    # En el año calendario, el producto sin mes cuenta normalmente.
    anual = informes.memoria(2026)
    assert {p["titulo"] for p in anual["productos"]} == {"Noviembre 2026", "Marzo 2026", "Sin mes 2026"}
    assert anual["por_confirmar"] == []
    rango = informes.memoria(desde="2026-03-01", hasta="2026-03-31")
    assert [p["titulo"] for p in rango["productos"]] == ["Marzo 2026"]


def test_rutas_de_memoria_por_periodo(client, emmy):
    _producto(emmy, "Sin mes 2026", 2026)
    _actividad(emmy, "Seminario de octubre", "2026-10-15")  # con actividades en el periodo
    login(client, "admin@test.cl", PASS_ADMIN)
    r = client.get("/informes/memoria?modo=periodo&periodo=2026&anio=2025")
    assert "Memoria del periodo 2026–2027" in r.text and "Por confirmar" in r.text
    assert "/informes/memoria/xlsx?periodo=2026" in r.text
    r = client.get("/informes/memoria?modo=rango&desde=2026-01-01&hasta=2026-06-30")
    assert "Memoria del 1 ene 2026 al 30 jun 2026" in r.text
    # Un rango inválido vuelve al año calendario.
    assert "Memoria anual" in client.get("/informes/memoria?modo=rango&desde=2026-06-30&hasta=2026-01-01").text
    r = client.get("/informes/memoria/xlsx?periodo=2026")
    assert r.headers["content-disposition"] == 'attachment; filename="memoria-2026-2027.xlsx"'
    assert "Sin mes (por confirmar)" in load_workbook(io.BytesIO(r.content)).sheetnames


# --- Plan de actividades ---------------------------------------------------------------------------------------

@pytest.fixture
def aceptada(client, emmy, monkeypatch):
    """Propuesta de Emmy aceptada para 2026–2027 (con dos actividades previstas)."""
    monkeypatch.setattr(inst, "hoy", lambda: date(2026, 9, 10))
    prop_id, _ = inst.crear_nuevo({"id": emmy, "es_investigador": True}, "Anillos")
    prop = inst.obtener(prop_id)
    inst.guardar(prop, {k: COMPLETA[k] for k in ("nombre", "descripcion", "objetivo_general", "metodologia",
                                                 "resultados_esperados", "bibliografia")},
                 COMPLETA["objetivo"], list(zip(COMPLETA["act_desc"], COMPLETA["act_mes"])))
    assert inst.enviar(inst.obtener(prop_id), {"id": emmy, "es_admin": False}) is None
    admin = usuarios.por_email("admin@test.cl")
    assert inst.decidir(inst.obtener(prop_id), "aceptada", "", {**admin, "es_admin": True}) is None
    return inst.obtener(prop_id)


def test_plan_previstas_y_realizadas(client, emmy, aceptada, monkeypatch):
    pid = aceptada["proyecto_id"]
    seminario = _actividad(emmy, "Seminario de lanzamiento", "2026-10-20")
    fuera = _actividad(emmy, "Algo de 2025", "2025-05-01")
    for a in (seminario, fuera):
        proyectos.sincronizar("actividad", a, [pid], {"id": emmy})  # Emmy es responsable: aprobado
    monkeypatch.setattr(inst, "hoy", lambda: date(2027, 4, 15))

    plan = inst.plan(pid)
    assert [(f["descripcion"], f["estado"]) for f in plan[0]["previstas"]] == [
        ("Seminario de lanzamiento", "atrasada"), ("Visita a Gotinga", "atrasada")]
    # Solo se enlazan actividades aprobadas del mismo periodo.
    assert [a["id"] for a in inst.actividades_para_enlazar(pid, 2026)] == [seminario]

    login(client, "inv@test.cl", PASS_TEMP)
    prevista = plan[0]["previstas"][0]["id"]
    r = client.post(f"/proyectos/{pid}/plan/{prevista}", headers=HX, data={"actividad_id": str(fuera)})
    assert inst.plan(pid)[0]["realizadas"] == 0
    r = client.post(f"/proyectos/{pid}/plan/{prevista}", headers=HX, data={"actividad_id": str(seminario)})
    assert "1 de 2 realizadas" in r.text
    assert inst.plan(pid)[0]["previstas"][0]["estado"] == "realizada"

    informe = client.get(f"/proyectos/{pid}/informe?desde=2026&hasta=2027").text
    assert "Plan de actividades" in informe and "1 de 2 actividades previstas realizadas" in informe
    libro = load_workbook(io.BytesIO(client.get(f"/proyectos/{pid}/informe/xlsx?desde=2026&hasta=2027").content))
    assert "Plan" in libro.sheetnames
    # Desmarcar.
    client.post(f"/proyectos/{pid}/plan/{prevista}", headers=HX, data={"actividad_id": ""})
    assert inst.plan(pid)[0]["realizadas"] == 0


def test_plan_solo_en_institucionales_y_solo_responsables(client, emmy, aceptada):
    pid = aceptada["proyecto_id"]
    alan = usuarios.crear("alan@test.cl", "Alan", "Turing", "Clave-colega-1", debe_cambiar_pass=False)["id"]
    proyectos.agregar_participante(pid, f"i:{alan}", "co_investigador", False)
    login(client, "alan@test.cl", "Clave-colega-1")
    assert "Plan" in client.get(f"/proyectos/{pid}").text
    prevista = inst.plan(pid)[0]["previstas"][0]["id"]
    assert client.post(f"/proyectos/{pid}/plan/{prevista}", headers=HX, data={"actividad_id": ""}).status_code == 403
    indiv = proyectos.crear(dict(nombre="Individual", estado="vigente"), alan, alan)
    assert "?tab=plan" not in client.get(f"/proyectos/{indiv}").text


# --- Exportación de la convocatoria ------------------------------------------------------------------------------

def test_exportar_convocatoria(client, emmy, aceptada):
    login(client, "inv@test.cl", PASS_TEMP)
    assert client.get("/admin/propuestas/2026/xlsx").status_code == 403
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "/admin/propuestas/2026/xlsx" in client.get("/admin/propuestas?periodo=2026").text
    r = client.get("/admin/propuestas/2026/xlsx")
    assert r.headers["content-disposition"] == 'attachment; filename="propuestas-2026-2027.xlsx"'
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Propuestas", "Textos", "Objetivos", "Actividades previstas"]
    fila = [c.value for c in libro["Propuestas"][2]]
    assert fila[:4] == ["Anillos y autómatas", "Proyecto nuevo", "Emmy Noether", "Aceptada"]
    assert libro["Objetivos"].max_row == 3 and libro["Actividades previstas"]["B2"].value == "oct 2026"
