"""Estudiantes en la investigación (IAC-CINDA 2.3.2, 2.3.4 y 3.6.1 d)."""
import io

import pytest
from openpyxl import load_workbook

from app import acreditacion, estudiantes, informes, personas, productos, proyectos, revistas, usuarios, vinculos
from app.db import get_db
from tests.conftest import HX, PASS_ADMIN, PASS_TEMP, login


def _tipo(nombre):
    return str(get_db().q("SELECT id FROM tipo_producto WHERE nombre = ?", [nombre])[0]["id"])


@pytest.fixture
def emmy(app):
    uid = usuarios.por_email("inv@test.cl")["id"]
    usuarios.cambiar_password(uid, PASS_TEMP, temporal=False)
    vinculos.agregar(uid, "titular", "2015-01-01")
    return uid


@pytest.fixture
def maria(emmy):
    mid = personas.crear_externo("María", "Quispe", "UMSA", "Bolivia", emmy)["id"]
    assert estudiantes.agregar(mid, {"nivel": "licenciatura", "desde": "2021", "hasta": "2024", "titulado": "1"}, emmy) == []
    assert estudiantes.agregar(mid, {"nivel": "maestria", "desde": "2025"}, emmy) == []
    return mid


def test_periodos_y_estados(emmy, maria):
    assert estudiantes.estudiante_en(maria, 2023)["nivel"] == "licenciatura"
    assert estudiantes.estudiante_en(maria, 2026)["nivel"] == "maestria"
    assert estudiantes.graduado_en(maria, 2025)["nivel"] == "licenciatura"
    assert estudiantes.graduado_en(maria, 2024) is None  # egresó ese año: aún no «graduada» ese año
    assert estudiantes.etiqueta(maria, 2026) == "Estudiante de maestría"
    assert estudiantes.periodos(maria)[0]["programa"] == "Licenciatura en Matemática (Universidad Mayor de San Andrés)"
    assert "se superpone" in estudiantes.agregar(maria, {"nivel": "maestria", "desde": "2026"}, emmy)[0]
    assert "anterior al ingreso" in estudiantes.agregar(maria, {"nivel": "doctorado", "desde": "2030", "hasta": "2029"}, emmy)[0]
    assert "año de egreso" in estudiantes.agregar(maria, {"nivel": "doctorado", "desde": "2030", "titulado": "1"}, emmy)[0]


def test_registrar_desde_la_pagina(client, emmy):
    login(client, "inv@test.cl", PASS_TEMP)
    assert "Registrar estudiante" in client.get("/estudiantes").text
    r = client.post("/estudiantes", headers=HX, data={"nombres": "Juan", "apellidos": "Mamani", "nivel": "licenciatura",
                                                      "desde": "2023", "del_programa": "1"})
    eid = int(r.headers["hx-redirect"].rsplit("/", 1)[1])
    assert estudiantes.periodos(eid)[0]["desde"] == 2023
    # Una persona externa ya registrada (p. ej., coautor) se elige por su id.
    otra = personas.crear_externo("Rosa", "Choque", None, None, emmy)["id"]
    r = client.post("/estudiantes", headers=HX, data={"externo_id": str(otra), "nivel": "doctorado", "desde": "2024"})
    assert r.headers["hx-redirect"] == f"/estudiantes/{otra}" and estudiantes.periodos(otra)[0]["nivel"] == "doctorado"
    r = client.post("/estudiantes", headers=HX, data={"nivel": "licenciatura", "desde": "2023"})
    assert "Elija a la persona" in r.text
    assert "Rosa" in client.get("/estudiantes/buscar-persona?persona_q=choque", headers=HX).text
    assert "Juan" in client.get("/estudiantes/lista?nivel=licenciatura", headers=HX).text
    assert "Juan" not in client.get("/estudiantes/lista?nivel=doctorado", headers=HX).text


def test_ficha_y_periodos(client, emmy, maria):
    login(client, "inv@test.cl", PASS_TEMP)
    r = client.get(f"/estudiantes/{maria}")
    assert "María Quispe" in r.text and "Maestría" in r.text
    r = client.post(f"/estudiantes/{maria}/periodo", headers=HX, data={"nivel": "doctorado", "desde": "2020", "hasta": "2019"})
    assert "anterior al ingreso" in r.text
    pid = estudiantes.periodos(maria)[1]["id"]
    client.put(f"/estudiantes/periodo/{pid}", headers=HX, data={"nivel": "maestria", "desde": "2025", "hasta": "2026",
                                                                "titulado": "1", "del_programa": "0"})
    p = estudiantes.obtener_periodo(pid)
    assert (p["hasta"], p["titulado"], p["del_programa"]) == (2026, 1, 0)
    client.delete(f"/estudiantes/periodo/{pid}", headers=HX)
    assert len(estudiantes.periodos(maria)) == 1


def test_distintivo_en_el_selector(client, emmy, maria):
    login(client, "inv@test.cl", PASS_TEMP)
    assert "Estudiante de maestría" in client.get(f"/personas/fila?campo=autor&ref=e:{maria}", headers=HX).text


def test_indicadores_y_anexo(client, emmy, maria):
    admin = usuarios.por_email("admin@test.cl")["id"]
    rid = revistas.crear({"titulo": "Revista X", "pais": "Bolivia", "arbitrada": 1}, admin, verificada=True)
    def producto(titulo, tipo, anio, autores, estado="publicado", revision="arbitrado", revista=None):
        d, errores = productos.preparar(dict(tipo_id=_tipo(tipo), titulo=titulo, anio=str(anio), estado_publicacion=estado,
                                             revision=revision, revista_id=str(revista or "")), autores, [])
        assert not errores, errores
        return productos.guardar(d, autores, emmy, filiaciones=[1 if a.startswith("i:") else None for a in autores])
    producto("Tesis de licenciatura", "Tesis dirigida", 2024, [f"e:{maria}", f"i:{emmy}"], revision="sin_arbitraje")
    producto("Tesis de maestría", "Tesis dirigida", 2026, [f"e:{maria}", f"i:{emmy}"], estado="en_preparacion",
             revision="sin_arbitraje")
    producto("Artículo con estudiante", "Artículo en revista", 2023, [f"i:{emmy}", f"e:{maria}"], revista=rid)
    producto("Artículo con graduada", "Artículo en revista", 2025, [f"i:{emmy}", f"e:{maria}"], revista=rid)
    # 2025: María ya es graduada de licenciatura y estudiante de maestría: cuenta en ambos.
    pid = proyectos.crear(dict(nombre="Proyecto con tesista", estado="vigente", fecha_inicio="2023-01-01",
                               fecha_fin="2024-12-31"), emmy, emmy)
    proyectos.agregar_participante(pid, f"e:{maria}", "tesista", False)

    e = acreditacion.anexo(2022, 2026)["estudiantes"]
    assert [(t["producto"]["titulo"], t["nivel"], t["estado"]) for t in e["tesis"]] == [
        ("Tesis de licenciatura", "licenciatura", "defendida"), ("Tesis de maestría", "maestria", "en curso")]
    assert e["tesis_por_nivel"] == {("Licenciatura", "defendida"): 1, ("Maestría", "en curso"): 1}
    assert {p["titulo"] for p in e["productos"]} == {"Artículo con estudiante", "Artículo con graduada"}
    assert [p["titulo"] for p in e["graduados"]] == ["Artículo con graduada"]
    assert [p["nombre"] for p in e["proyectos"]] == ["Proyecto con tesista"]
    assert [g for g in informes.cv(emmy, 2022, 2026)["grupos"]["tesis"]]  # el CV las agrupa

    login(client, "admin@test.cl", PASS_ADMIN)
    r = client.get("/informes/acreditacion?desde=2022&hasta=2026")
    assert "Estudiantes en la investigación" in r.text and "Tesis de licenciatura" in r.text
    assert "Participación de estudiantes en tesis" not in r.text  # ya no figura como pendiente
    libro = load_workbook(io.BytesIO(client.get("/informes/acreditacion/xlsx?desde=2022&hasta=2026").content))
    assert libro["Tesis dirigidas"].max_row == 3 and "Productos con estudiantes" in libro.sheetnames


def test_solo_con_director_vinculado(emmy, maria):
    """Una tesis sin director del instituto vinculado en la fecha no cuenta."""
    otro = usuarios.crear("x@test.cl", "Sin", "Vínculo", PASS_TEMP)["id"]
    d, _ = productos.preparar(dict(tipo_id=_tipo("Tesis dirigida"), titulo="Ajena", anio="2024", estado_publicacion="publicado",
                                   revision="sin_arbitraje"), [f"e:{maria}", f"i:{otro}"], [])
    productos.guardar(d, [f"e:{maria}", f"i:{otro}"], otro)
    assert estudiantes.tesis(2020, 2026) == []
