"""Tutoría, directorio de egresados, mentoría y anexo de seguimiento estudiantil (PLAN §49 y §50, etapa 5)."""
import io
import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from app.acreditacion import seguimiento_estudiantil as seg
from app.comun import estudiantes, notificaciones, permisos, tokens, unidades, usuarios
from app.comun.db import get_db
from app.egresados import directorio, egresados
from app.estudiantil import encuestas, importar, tutoria
from tests.conftest import HX, PASS_ADMIN, PASS_TEMP, login
from tests.test_cohortes import ENCABEZADO
from tests.test_portal_estudiante import _aceptar, _ingresar

CLAVE = "Clave-roles-1"
BASE = "/docencia/tutoria"


def _cargar(*filas):
    filas, errores = importar.leer((ENCABEZADO + "\n".join(filas) + "\n").encode())
    assert errores == [] and not [f["errores"] for f in filas if f["errores"]], [f["errores"] for f in filas]
    importar.aplicar(filas, None)


def _uid(email):
    return usuarios.por_email(email)["id"]


def _externo(ru):
    return get_db().q("SELECT externo_id FROM estudiante WHERE registro_universitario = ?", [ru])[0]["externo_id"]


def _cliente(app, email, clave=CLAVE):
    c = TestClient(app)
    assert login(c, email, clave).headers["hx-redirect"] == "/"
    return c


@pytest.fixture
def mundo(app):
    """Dos tutores (con vínculo), un comité, tres estudiantes que cursan (dos de la cohorte 2022-I), una titulada y una retirada."""
    for email, nombre in (("tutor1@test.cl", "Tutor1"), ("tutor2@test.cl", "Tutor2")):
        usuarios.crear(email, nombre, "Docente", CLAVE, debe_cambiar_pass=False, categoria="titular", fecha_ingreso="2015-01-01")
    c = usuarios.crear("comite@test.cl", "Comité", "Prueba", CLAVE, debe_cambiar_pass=False, es_investigador=False)
    permisos.asignar(c["id"], "comite_acreditacion", unidades.id_de("UMSA"))
    hace = (date.today() - timedelta(days=400)).isoformat()
    _cargar("8000001,Ana,Alumna,ana@correo.bo,mujer,,2022-I,,,,,,cursando,,,,,",
            "8000002,Beto,Alumno,beto@correo.bo,hombre,,2022-I,,,,,,cursando,,,,,",
            "8000003,Cora,Alumna,cora@correo.bo,mujer,,2023-II,,,,,,cursando,,,,,",
            f"8000004,Dina,Titulada,dina@correo.bo,mujer,,2018-I,,,,,,titulado,,{hace},{hace},tesis,",
            "8000005,Eva,Retirada,eva@correo.bo,mujer,,2020-I,,,,,,retirado,2021-I,,,,academico")


@pytest.fixture
def direccion(app, mundo):
    permisos.asignar(_uid("admin@test.cl"), "direccion_carrera", unidades.id_de("CMAT"))
    return _cliente(app, "admin@test.cl", PASS_ADMIN)


# --- Tutoría: asignación -----------------------------------------------------------------------------------------

def test_solo_los_docentes_con_vinculo_pueden_ser_tutores(mundo):
    assert {t["email"] for t in tutoria.tutores_elegibles()} == {"tutor1@test.cl", "tutor2@test.cl"}
    ana = _externo("8000001")
    assert tutoria.asignar(ana, _uid("inv@test.cl"), None) == ["Elija un docente con vínculo vigente."]
    assert tutoria.asignar(_externo("8000004"), _uid("tutor1@test.cl"), None) == ["Solo se asigna tutor a quien cursa la licenciatura."]
    assert tutoria.asignar(_externo("8000005"), _uid("tutor1@test.cl"), None) == ["Solo se asigna tutor a quien cursa la licenciatura."]
    assert tutoria.asignar(ana, _uid("tutor1@test.cl"), None) == []


def test_cambiar_de_tutor_cierra_la_asignacion_anterior(mundo):
    ana = _externo("8000001")
    tutoria.asignar(ana, _uid("tutor1@test.cl"), None)
    tutoria.asignar(ana, _uid("tutor2@test.cl"), None)
    filas = get_db().q("SELECT tutor_id, hasta FROM tutoria_asignacion WHERE externo_id = ? ORDER BY id", [ana])
    assert filas[0]["hasta"] is not None and filas[1]["hasta"] is None
    assert estudiantes.tutor_de(ana)["email"] == "tutor2@test.cl" and not estudiantes.es_tutor(_uid("tutor1@test.cl"))


def test_asignar_a_una_cohorte_solo_a_quien_no_tiene_tutor(mundo):
    tutoria.asignar(_externo("8000002"), _uid("tutor2@test.cl"), None)
    assert tutoria.asignar_cohorte(2022, "I", _uid("tutor1@test.cl"), None) == 1
    assert estudiantes.tutor_de(_externo("8000001"))["email"] == "tutor1@test.cl"
    assert estudiantes.tutor_de(_externo("8000002"))["email"] == "tutor2@test.cl"          # no se le quitó
    assert estudiantes.tutor_de(_externo("8000003")) is None
    assert notificaciones.recientes(_uid("tutor1@test.cl"))[0]["asunto"] == "Le asignaron estudiantes como tutor"


# --- Tutoría: reuniones --------------------------------------------------------------------------------------------

REUNION = {"fecha": date.today().isoformat(), "tema": "academico", "resumen": "Dificultades en Álgebra", "derivacion": "bienestar",
           "acuerdos": "Volver en un mes"}


def test_solo_el_tutor_registra_y_ve_sus_reuniones(mundo):
    ana, t1, t2 = _externo("8000001"), _uid("tutor1@test.cl"), _uid("tutor2@test.cl")
    tutoria.asignar(ana, t1, None)
    assert tutoria.registrar(ana, t2, REUNION) == ["Ese estudiante no está entre sus tutorados."]
    assert tutoria.registrar(ana, t1, {**REUNION, "tema": "", "fecha": "2099-01-01"}) == ["La fecha no puede ser futura.", "Elija el tema."]
    assert tutoria.registrar(ana, t1, REUNION) == []
    assert [r["resumen"] for r in tutoria.reuniones(ana, t1)] == ["Dificultades en Álgebra"]
    # Con otro tutor, no ve las notas del anterior; la Dirección (sin filtro) ve todas.
    tutoria.asignar(ana, t2, None)
    assert tutoria.reuniones(ana, t2) == [] and len(tutoria.reuniones(ana)) == 1
    rid = tutoria.reuniones(ana)[0]["id"]
    assert tutoria.editar(rid, t2, REUNION) == ["Reunión no encontrada."] and not tutoria.eliminar(rid, t2)
    assert tutoria.editar(rid, t1, {**REUNION, "resumen": "Corregido"}) == []
    assert tutoria.obtener(rid)["resumen"] == "Corregido" and tutoria.eliminar(rid, t1)


def test_pantallas_del_tutor(app, mundo):
    ana, beto = _externo("8000001"), _externo("8000002")
    tutoria.asignar(ana, _uid("tutor1@test.cl"), None)
    tutoria.asignar(beto, _uid("tutor2@test.cl"), None)
    c = _cliente(app, "tutor1@test.cl")
    r = c.get("/tutoria")
    assert r.status_code == 200 and "Alumna, Ana" in r.text and "Alumno, Beto" not in r.text and 'href="/tutoria"' in r.text
    assert c.post(f"/tutoria/{ana}/reuniones", data=REUNION, headers=HX).headers["hx-redirect"] == f"/tutoria/{ana}"
    r = c.get(f"/tutoria/{ana}")
    assert "Dificultades en Álgebra" in r.text and "Registrar reunión" in r.text and "Confidencial" in r.text
    assert c.get(f"/tutoria/{beto}").status_code == 403                                       # no es su tutorado
    assert c.get(f"/tutoria/{beto}/reuniones/nueva", headers=HX).status_code == 403
    assert c.post(f"/tutoria/{beto}/reuniones", data=REUNION, headers=HX).status_code == 403
    rid = tutoria.reuniones(ana)[0]["id"]
    assert c.post(f"/tutoria/reuniones/{rid}", data={**REUNION, "resumen": "Cambiado"}, headers=HX).headers["hx-redirect"] == f"/tutoria/{ana}"
    assert tutoria.obtener(rid)["resumen"] == "Cambiado"
    otro = _cliente(app, "tutor2@test.cl")
    assert otro.post(f"/tutoria/reuniones/{rid}", data=REUNION, headers=HX).status_code == 403
    assert otro.post(f"/tutoria/reuniones/{rid}/eliminar", headers=HX).status_code == 403
    assert tutoria.obtener(rid) is not None
    assert c.post(f"/tutoria/reuniones/{rid}/eliminar", headers=HX).headers["hx-redirect"] == f"/tutoria/{ana}"
    assert tutoria.obtener(rid) is None


def test_quien_no_es_tutor_no_ve_tutoria(app, mundo):
    usuarios.cambiar_password(_uid("inv@test.cl"), PASS_TEMP, temporal=False)
    c = _cliente(app, "inv@test.cl", PASS_TEMP)
    assert c.get("/tutoria").status_code == 403 and 'href="/tutoria"' not in c.get("/").text


def test_la_direccion_ve_todas_las_reuniones_y_queda_en_la_auditoria(direccion):
    ana = _externo("8000001")
    tutoria.asignar(ana, _uid("tutor1@test.cl"), None)
    tutoria.registrar(ana, _uid("tutor1@test.cl"), REUNION)
    r = direccion.get(f"/tutoria/{ana}")
    assert r.status_code == 200 and "Dificultades en Álgebra" in r.text and "Registrar reunión" not in r.text
    assert "queda en la auditoría" in r.text
    assert get_db().q("SELECT COUNT(*) n FROM auditoria WHERE accion = 'ver_tutorias'")[0]["n"] == 1


def test_pantalla_de_la_direccion_y_asignacion(direccion):
    r = direccion.get(BASE)
    assert r.status_code == 200 and "Estudiantes que cursan (3)" in r.text and "Alumna, Ana" in r.text
    assert "Titulada" not in r.text and "Retirada" not in r.text
    ana = _externo("8000001")
    direccion.post(f"{BASE}/{ana}/asignar", data={"tutor_id": str(_uid("tutor1@test.cl"))}, headers=HX)
    assert estudiantes.tutor_de(ana)["email"] == "tutor1@test.cl"
    direccion.post(f"{BASE}/asignar-cohorte", data={"cohorte": "2022-I", "tutor_id": str(_uid("tutor2@test.cl"))}, headers=HX)
    assert estudiantes.tutor_de(_externo("8000002"))["email"] == "tutor2@test.cl" and estudiantes.tutor_de(ana)["email"] == "tutor1@test.cl"
    assert "Solo sin tutor" in direccion.get(f"{BASE}?solo=sin_tutor").text
    # Datos inválidos no rompen.
    direccion.post(f"{BASE}/asignar-cohorte", data={"cohorte": "xx", "tutor_id": "1"}, headers=HX)
    direccion.post(f"{BASE}/{ana}/asignar", data={"tutor_id": str(_uid("inv@test.cl"))}, headers=HX)
    assert estudiantes.tutor_de(ana)["email"] == "tutor1@test.cl"


def test_el_comite_ve_cifras_de_tutoria_pero_no_nombres_ni_reuniones(app, mundo):
    ana = _externo("8000001")
    tutoria.asignar(ana, _uid("tutor1@test.cl"), None)
    tutoria.registrar(ana, _uid("tutor1@test.cl"), REUNION)
    c = _cliente(app, "comite@test.cl")
    r = c.get(BASE)
    assert r.status_code == 200 and "Reuniones registradas" in r.text and "Alumna" not in r.text and "Dificultades" not in r.text
    assert c.get(f"/tutoria/{ana}").status_code == 403
    assert c.post(f"{BASE}/{ana}/asignar", data={"tutor_id": str(_uid("tutor2@test.cl"))}, headers=HX).status_code == 403
    ind = tutoria.indicadores()
    assert ind["con_tutor"] == 1 and ind["cursando"] == 3 and ind["reuniones"] == 1 and ind["derivaciones"]["Bienestar universitario"] == 1


def test_el_estudiante_ve_su_tutor_pero_no_las_notas(app, mundo, buzon):
    ana = _externo("8000001")
    estudiantes.crear_cuentas(None)
    tutoria.asignar(ana, _uid("tutor1@test.cl"), None)
    tutoria.registrar(ana, _uid("tutor1@test.cl"), REUNION)
    c = TestClient(app)
    _ingresar(c, "ana@correo.bo", buzon)
    _aceptar(c)
    r = c.get("/estudiante")
    assert "Su tutor(a)" in r.text and "Tutor1 Docente" in r.text and "Dificultades" not in r.text
    for ruta in ("/tutoria", f"/tutoria/{ana}", BASE):
        assert c.get(ruta, follow_redirects=False).status_code == 303                       # fuera del portal del estudiante


# --- Directorio ------------------------------------------------------------------------------------------------

DINA = {"ocupacion": "Analista de datos", "institucion": "Banco", "ciudad_pais": "La Paz, Bolivia", "areas": "Estadística",
        "ofrece": ["charla", "mentoria"], "detalle": "Puedo dar charlas", "visible": "on", "acepto": "on"}


def test_invitar_crea_la_cuenta_y_manda_un_enlace_de_siete_dias(mundo, buzon):
    dina, ana = _externo("8000004"), _externo("8000001")
    assert directorio.invitar(ana, None) == ["Solo se invita a egresados titulados."]
    assert directorio.invitar(dina, None) == []
    u = estudiantes.usuario_de(dina)
    assert u["email"] == "dina@correo.bo" and directorio.estado_directorio(dina) == "invitado"
    assert "/acceso/" in buzon[-1].get_content() and "opcional" in buzon[-1].get_content()
    token = buzon[-1].get_content().split("/acceso/")[1].split()[0]
    t = get_db().q("SELECT expira_en, creado_en FROM token_acceso WHERE usuario_id = ?", [u["id"]])[0]
    assert t["expira_en"] > (date.today() + timedelta(days=6)).isoformat() and tokens.vigente(token, ("acceso",))
    assert directorio.invitar(dina, None) == []                                          # reenviar no falla ni duplica la cuenta
    assert get_db().q("SELECT COUNT(*) n FROM cuenta_estudiante WHERE externo_id = ?", [dina])[0]["n"] == 1


def test_invitar_con_correo_en_uso_o_faltante(app):
    hace = (date.today() - timedelta(days=400)).isoformat()
    _cargar(f"8100001,Uno,Titulado,inv@test.cl,mujer,,2018-I,,,,,,titulado,,{hace},{hace},tesis,",
            f"8100002,Dos,Titulado,,mujer,,2018-I,,,,,,titulado,,{hace},{hace},tesis,")
    assert directorio.invitar(_externo("8100001"), None) == ["El correo ya lo usa otra cuenta: corríjalo antes de invitar."]
    assert directorio.invitar(_externo("8100002"), None) == ["El egresado no tiene correo registrado."]
    assert directorio.invitar_pendientes(None) == {"invitados": 0, "con_problemas": 1}


def test_el_perfil_exige_consentimiento_y_ocupacion(mundo):
    dina = _externo("8000004")
    assert directorio.guardar_perfil(_externo("8000001"), DINA) == ["Solo los egresados titulados figuran en el directorio."]
    assert directorio.guardar_perfil(dina, {**DINA, "acepto": ""})[0].startswith("Para figurar en el directorio debe aceptar")
    assert directorio.guardar_perfil(dina, {**DINA, "ocupacion": ""}) == ["Cuente a qué se dedica para figurar en el directorio."]
    assert directorio.guardar_perfil(dina, DINA) == []
    p = directorio.perfil(dina)
    assert p["visible"] == 1 and p["consentimiento_en"] and p["ofrece"] == ["charla", "mentoria"]
    assert [x["nombres"] for x in directorio.visibles()] == ["Dina"]
    assert directorio.visibles("banco") and not directorio.visibles("zoología") and directorio.visibles(ofrece="mentoria")
    assert not directorio.visibles(ofrece="empleo")
    assert directorio.guardar_perfil(dina, {**DINA, "visible": ""}) == [] and directorio.visibles() == []     # puede retirarse
    assert directorio.perfil(dina)["consentimiento_en"]                                                       # conserva cuándo aceptó


def _con_directorio(buzon):
    """Dina figura en el directorio; Ana y Beto tienen cuenta."""
    estudiantes.crear_cuentas(None)
    dina = _externo("8000004")
    directorio.invitar(dina, None)
    estudiantes.aceptar_consentimiento(dina)
    assert directorio.guardar_perfil(dina, DINA) == []
    return dina


def test_solicitud_de_mentoria_comparte_el_correo_solo_al_aceptar(mundo, buzon):
    dina, ana = _con_directorio(buzon), _externo("8000001")
    assert directorio.solicitar(ana, dina, "", "Hola") == ["Elija qué necesita."]
    assert directorio.solicitar(ana, dina, "charla", " ") == ["Cuéntele brevemente qué necesita."]
    assert directorio.solicitar(ana, _externo("8000003"), "charla", "Hola") == ["Ese egresado no figura en el directorio."]
    assert directorio.solicitar(ana, dina, "charla", "¿Podría dar una charla de datos?") == []
    assert directorio.solicitar(ana, dina, "charla", "otra vez") == ["Ya le envió una solicitud a esta persona y aún no responde."]
    aviso = notificaciones.recientes(estudiantes.usuario_de(dina)["id"])[0]
    assert aviso["asunto"] == "Un estudiante le pidió su ayuda" and "ana@correo.bo" not in aviso["texto"] and "Ana Alumna" in aviso["texto"]
    sid = directorio.recibidas(dina)[0]["id"]
    assert not directorio.responder(sid, _externo("8000002"), True)                      # solo la destinataria responde
    assert directorio.responder(sid, dina, True) and not directorio.responder(sid, dina, False)
    aviso = notificaciones.recientes(estudiantes.usuario_de(ana)["id"])[0]
    assert aviso["asunto"] == "Aceptaron su solicitud" and "dina@correo.bo" in aviso["texto"]
    assert directorio.enviadas(ana)[0]["estado"] == "aceptada" and directorio.indicadores()["aceptadas"] == 1


def test_limite_de_solicitudes_pendientes(mundo, buzon):
    dina, ana = _con_directorio(buzon), _externo("8000001")
    hace = (date.today() - timedelta(days=400)).isoformat()
    _cargar(*[f"82000{i},Gr{i},Titulado,gr{i}@correo.bo,mujer,,2018-I,,,,,,titulado,,{hace},{hace},tesis," for i in range(3)])
    for i in range(3):
        e = _externo(f"82000{i}")
        estudiantes.crear_cuenta(e, None)
        assert directorio.guardar_perfil(e, DINA) == []
    destinos = [_externo(f"82000{i}") for i in range(3)] + [dina]
    for a in destinos[:3]:
        assert directorio.solicitar(ana, a, "mentoria", "Hola") == []
    assert directorio.solicitar(ana, destinos[3], "mentoria", "Hola") == [
        f"Ya tiene {directorio.MAX_PENDIENTES} solicitudes pendientes: espere las respuestas."]
    assert directorio.solicitar(ana, ana, "mentoria", "Hola") == ["No puede pedirse algo a sí mismo."]


def test_pantallas_del_directorio(app, mundo, buzon):
    dina = _con_directorio(buzon)
    c = TestClient(app)
    _ingresar(c, "ana@correo.bo", buzon)
    _aceptar(c)
    r = c.get("/estudiante")
    assert "Directorio de egresados" in r.text and "Mi perfil de egresado" not in r.text
    r = c.get("/estudiante/directorio")
    assert "Dina Titulada" in r.text and "Analista de datos" in r.text and "Pedir ayuda" in r.text and "dina@correo.bo" not in r.text
    assert "Dina Titulada" not in c.get("/estudiante/directorio?q=zoologia").text
    assert c.get("/estudiante/perfil-egresado").status_code == 403                      # los estudiantes no tienen perfil de egresado
    assert "Pedir ayuda a un egresado" in c.get(f"/estudiante/directorio/{dina}/solicitar", headers=HX).text
    r = c.post(f"/estudiante/directorio/{dina}/solicitar", data={"tipo": "charla", "mensaje": ""}, headers=HX)
    assert "Cuéntele brevemente" in r.text
    r = c.post(f"/estudiante/directorio/{dina}/solicitar", data={"tipo": "charla", "mensaje": "Una charla, por favor"}, headers=HX)
    assert r.headers["hx-redirect"] == "/estudiante/directorio"
    assert "Mis solicitudes" in c.get("/estudiante/directorio").text
    # La egresada ve y responde.
    d = TestClient(app)
    _ingresar(d, "dina@correo.bo", buzon)
    r = d.get("/estudiante/perfil-egresado")
    assert r.status_code == 200 and "Una charla, por favor" in r.text and "Alumna" in r.text and "ana@correo.bo" not in r.text
    sid = directorio.recibidas(dina)[0]["id"]
    assert d.post(f"/estudiante/solicitudes/{sid}/aceptar", headers=HX).headers["hx-redirect"] == "/estudiante/perfil-egresado"
    assert "ana@correo.bo" in d.get("/estudiante/perfil-egresado").text
    assert "dina@correo.bo" in c.get("/estudiante/directorio").text
    c.post(f"/estudiante/solicitudes/{sid}/rechazar", headers=HX)                       # un estudiante no responde por la egresada
    assert directorio.recibidas(dina)[0]["estado"] == "aceptada"


def test_el_egresado_edita_su_perfil_desde_su_cuenta(app, mundo, buzon):
    dina = _externo("8000004")
    directorio.invitar(dina, None)
    d = TestClient(app)
    _ingresar(d, "dina@correo.bo", buzon)
    r = d.get("/estudiante")
    assert "Antes de empezar" in r.text                                                  # primero el consentimiento general
    _aceptar(d)
    assert "Mi perfil de egresado" in d.get("/estudiante").text
    r = d.post("/estudiante/perfil-egresado", data={"ocupacion": "", "visible": "on"})
    assert "Para figurar en el directorio debe aceptar" in r.text and directorio.perfil(dina)["visible"] == 0
    r = d.post("/estudiante/perfil-egresado", data={"ocupacion": "Docente", "visible": "on", "acepto": "on", "ofrece": ["charla", "empleo", "xx"]})
    assert r.status_code == 200 and directorio.perfil(dina)["ofrece"] == ["charla", "empleo"] and directorio.perfil(dina)["visible"] == 1
    assert [x["nombres"] for x in directorio.visibles()] == ["Dina"]


def test_directorio_del_personal(app, direccion, buzon):
    dina = _externo("8000004")
    r = direccion.get("/docencia/egresados")
    assert "Invitar" in r.text and "Invitar al directorio a quienes faltan" in r.text
    direccion.post(f"/docencia/egresados/directorio/invitar/{dina}", headers=HX)
    assert directorio.estado_directorio(dina) == "invitado" and "Invitado(a)" in direccion.get("/docencia/egresados").text
    assert direccion.get("/docencia/egresados/directorio").status_code == 200
    estudiantes.aceptar_consentimiento(dina)
    directorio.guardar_perfil(dina, DINA)
    assert "Dina Titulada" in direccion.get("/docencia/egresados/directorio").text and "Figura" in direccion.get("/docencia/egresados").text
    c = _cliente(app, "comite@test.cl")
    assert c.get("/docencia/egresados/directorio").status_code == 403
    assert c.post(f"/docencia/egresados/directorio/invitar/{dina}", headers=HX).status_code == 403
    assert c.post("/docencia/egresados/directorio/invitar-pendientes", headers=HX).status_code == 403


# --- Anexo de seguimiento estudiantil -----------------------------------------------------------------------------

def _con_todo(buzon):
    estudiantes.crear_cuentas(None)
    tutoria.asignar(_externo("8000001"), _uid("tutor1@test.cl"), None)
    tutoria.registrar(_externo("8000001"), _uid("tutor1@test.cl"), REUNION)
    dina = _con_directorio(buzon)
    directorio.solicitar(_externo("8000001"), dina, "charla", "Hola")
    eid = encuestas.crear_desde_plantilla("satisfaccion", None)
    encuestas.abrir(eid, None)
    inv = get_db().q("SELECT token FROM invitacion_encuesta WHERE encuesta_id = ?", [eid])
    qs = encuestas.preguntas(eid)
    encuestas.responder(inv[0]["token"], {q["id"]: {"escala": ["4"], "si_no": ["si"], "texto": ["Bien"]}[q["tipo"]] for q in qs})
    egresados.activar_seguimiento(None)
    return eid


def test_el_anexo_solo_trae_cifras_y_reserva_las_encuestas_con_pocas_respuestas(app, mundo, buzon):
    _con_todo(buzon)
    i = seg.informe(2015, date.today().year)
    assert i["cohortes"]["totales"]["total"] == 5 and i["tutoria"]["con_tutor"] == 1 and i["directorio"]["visibles"] == 1
    assert i["encuestas"][0]["oculto"] and i["encuestas"][0]["promedio_escalas"] is None and i["encuestas"][0]["respondieron"] == 1
    hojas = seg.filas_xlsx(i)
    assert list(hojas) == ["Cohortes", "Actividad estudiantil", "Encuestas", "Tutoría", "Egresados", "Seguimiento de egresados",
                           "Directorio y mentoría"]
    assert "reservados" in hojas["Encuestas"][0]["Nota"]
    texto = json.dumps(hojas, ensure_ascii=False)
    for privado in ("Alumna", "Titulada", "Dificultades", "ana@correo.bo", "dina@correo.bo", "Hola"):
        assert privado not in texto, privado


def test_pantalla_y_xlsx_del_anexo(app, direccion, buzon):
    _con_todo(buzon)
    r = direccion.get("/informes/seguimiento-estudiantil")
    assert r.status_code == 200 and "Seguimiento estudiantil" in r.text and "reservado" in r.text and "Alumna" not in r.text
    assert "Seguimiento estudiantil" in direccion.get("/informes").text
    x = direccion.get("/informes/seguimiento-estudiantil/xlsx")
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    libro = load_workbook(io.BytesIO(x.content))
    assert "Tutoría" in libro.sheetnames and "Directorio y mentoría" in libro.sheetnames
    c = _cliente(app, "comite@test.cl")
    assert c.get("/informes/seguimiento-estudiantil").status_code == 200 and "Dificultades" not in c.get("/informes/seguimiento-estudiantil").text
    usuarios.cambiar_password(_uid("inv@test.cl"), PASS_TEMP, temporal=False)
    i = _cliente(app, "inv@test.cl", PASS_TEMP)
    assert i.get("/informes/seguimiento-estudiantil").status_code == 403 and i.get("/informes/seguimiento-estudiantil/xlsx").status_code == 403
