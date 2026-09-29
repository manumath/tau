"""Guía de Antecedentes de IAC-CINDA: plantilla oficial llenada sin alterar su estructura."""
import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.academico import gestiones, oferta
from app.acreditacion import guia_antecedentes as guia
from app.comun import academicos, usuarios, vinculos
from app.comun.db import get_db
from app.comun.xlsx_plantilla import Plantilla
from app.investigacion import proyectos
from tests.conftest import PASS_ADMIN, PASS_TEMP, login
from tests.test_acreditacion import datos  # noqa: F401  (fixture)

H = guia.H


def _libro(contenido: bytes):
    return load_workbook(io.BytesIO(contenido))


def _partes(contenido: bytes) -> set[str]:
    return set(zipfile.ZipFile(io.BytesIO(contenido)).namelist())


def test_plantilla_conserva_graficos_imagen_y_estilos():
    original = guia.PLANTILLA.read_bytes()
    x = Plantilla(original)
    x.escribir(H["articulos"], "C6", "Título con & y <símbolos>")
    x.escribir(H["articulos"], "C90", "fila nueva")          # más allá de las filas de la plantilla
    x.escribir(H["academicos"], "K7", 55)
    x.formula(H["egresados"], "L7", "SUM(B7:K7)/'4. Cohortes ciclo-semestre'!B8")
    x.escribir(H["extension"], "Q9", 5)                    # maestra de una fórmula compartida (Q9:Q57)
    x.quitar_vinculos_externos()
    nuevo = x.guardar()
    antes, despues = _partes(original), _partes(nuevo)
    # Gráficos, imagen, dibujos y comentarios intactos; sin cadena de cálculo ni vínculo externo.
    assert {p for p in antes if p.startswith(("xl/charts/", "xl/media/", "xl/drawings/", "xl/comments"))} <= despues
    assert "xl/calcChain.xml" not in despues and not any(p.startswith("xl/externalLinks/") for p in despues)
    assert b'fullCalcOnLoad="1"' in zipfile.ZipFile(io.BytesIO(nuevo)).read("xl/workbook.xml")
    libro = _libro(nuevo)
    assert libro[H["articulos"]]["C6"].value == "Título con & y <símbolos>"
    assert libro[H["articulos"]]["C90"].value == "fila nueva"
    assert libro[H["articulos"]]["C90"].style_id == libro[H["articulos"]]["C55"].style_id  # estilo de la fila anterior
    assert libro[H["academicos"]]["K7"].value == 55
    assert libro[H["egresados"]]["L7"].value == "=SUM(B7:K7)/'4. Cohortes ciclo-semestre'!B8"
    assert libro.sheetnames == _libro(original).sheetnames
    ext = libro[H["extension"]]
    assert ext["Q9"].value == 5 and ext["Q10"].value == "=O10+P10" and ext["Q57"].value == "=O57+P57"
    # En el XML, las celdas que dependían de Q9 llevan ahora su fórmula explícita.
    hoja = zipfile.ZipFile(io.BytesIO(nuevo)).read(Plantilla(nuevo)._hojas[H["extension"]]).decode()
    assert '<c r="Q10" s=' in hoja and "<f>O10+P10</f>" in hoja and 'si="' not in hoja.split('r="Q9"')[1][:400]


def test_ciclos():
    assert guia.ciclos_hasta(2026, "II")[0] == (2022, "I") and len(guia.ciclos_hasta(2026, "II")) == 10
    assert guia.ciclos_hasta(2026, "I", 3) == [(2025, "I"), (2025, "II"), (2026, "I")]


def _institucional_en_bolivianos():
    get_db().execute("UPDATE proyecto SET monto = 69600, moneda = 'BOB', tipo = 'institucional', "
                     "clase = 'investigacion' WHERE id = ?", [proyectos.todos()[0]["id"]])


@pytest.fixture
def guia_llena(datos):  # noqa: F811
    emmy, alan = datos["emmy"], datos["alan"]
    g = {x["tipo"]: x for x in gestiones.crear_anio(2026)}
    oid, _ = oferta.agregar(g["I"]["id"], "MAT-111", "A", "40", 1)
    assert oferta.asignar(oid, emmy, "titular") == []
    oid2, _ = oferta.agregar(g["I"]["id"], "MAT-112", "A", "25", 1)
    assert oferta.asignar(oid2, alan, "auxiliar") == []
    assert academicos.agregar_grado(emmy, {"nivel": "doctorado", "titulo": "Doctora en Matemáticas",
                                    "institucion": "Universidad de Erlangen", "pais": "Alemania", "anio": "2010"}) == []
    assert academicos.fijar_fecha_nacimiento(emmy, "1982-03-23") == []
    _institucional_en_bolivianos()
    contenido = guia.generar(ultimo=(2026, "II"), tipo_cambio=6.96)
    return _libro(contenido), datos


def test_guia_con_los_datos_del_sistema(guia_llena):
    libro, datos = guia_llena
    ident = libro[H["id"]]
    assert ident["B6"].value == "Licenciatura en Matemática" and ident["B7"].value == "Universidad Mayor de San Andrés"
    assert [ident[f"B{r}"].value for r in (30, 39)] == ["2022-I", "2026-II"]
    assert (ident["A19"].value, ident["B19"].value) == ("CIENCIAS_NATURALES", "1.1. MATEMÁTICAS")

    asig = libro[H["asignaturas"]]
    fila = next(r for r in range(7, 80) if (asig[f"C{r}"].value or "").endswith("(MAT-111)"))
    assert (asig[f"B{fila}"].value, asig[f"E{fila}"].value, asig[f"F{fila}"].value) == (1, 1, 40)

    acad = libro[H["academicos"]]
    filas = {acad[f"C{r}"].value: r for r in range(7, 20) if acad[f"C{r}"].value}
    r = filas["Noether"]
    assert acad[f"E{r}"].value == "Doctora en Matemáticas" and acad[f"F{r}"].value == "Universidad de Erlangen, Alemania"
    assert acad[f"K{r}"].value == 44 and acad[f"N{r}"].value == "Titular" and acad[f"L{r}"].value == 2019
    assert acad[f"V{r}"].value == 1 and "Algebra I (MAT-111) (2026-I)" in acad[f"W{r}"].value
    assert "Turing" not in filas  # el invitado terminó en 2024: no es académico al cierre

    jefes = libro[H["jefes"]]
    assert (jefes["C9"].value, jefes["I9"].value, jefes["K9"].value) == ("Turing", "Auxiliar de docencia", 1)

    car = libro[H["caracterizacion"]]
    assert car["AC8"].value == "Titular" and car["A18"].value == "2026-II"
    assert car["B18"].value == "=SUM(E18:H18)"                           # la plantilla cuenta por grado
    assert sum(car[f"{c}18"].value for c in "EFGHLMNOSTUV") == 1          # Emmy, titular, en un solo grado
    assert car["AC18"].value == 1 and car["Z18"].value == 1             # 44 años: entre 36 y 50

    proy = libro[H["proyectos"]]
    assert proy["E9"].value == "Proyecto IDH" and proy["AT9"].value == 10000.0   # 69 600 Bs / 6,96
    assert proy["AW9"].value == "=SUM(AT9:AU9)" and proy["AW30"].value == "=SUM(AT30:AU30)"

    art = libro[H["articulos"]]
    titulos = {art[f"C{r}"].value: r for r in range(6, 20) if art[f"C{r}"].value}
    assert {"Nacional", "Internacional con coautor", "Declarado"} <= set(titulos)
    assert "Antes de la ventana" not in titulos                               # 2019: fuera de los 10 ciclos
    r = titulos["Internacional con coautor"]
    assert art[f"E{r}"].value == "Hermann Weyl" and art[f"J{r}"].value == "internacional"
    assert "Scopus" in art[f"K{r}"].value
    assert libro[H["libros"]]["D6"].value == "Un libro"
    assert libro[H["patentes"]]["C6"].value == "Una patente"
    assert libro[H["desarrollo"]]["C10"].value in {"Congreso con charla", "Congreso de oyente", "Estancia en Santiago"}
    assert libro[H["profesional"]]["G8"].value.startswith("Asesoría al ministerio")
    assert libro[H["redes"]]["B7"].value == "Red andina" and libro[H["redes"]]["C7"].value == "Internacional"

    # Fórmulas de la plantilla corregidas y ciclos de las hojas de cohortes.
    assert libro[H["egresados"]]["L7"].value == "=SUM(B7:K7)/'4. Cohortes ciclo-semestre'!B8"
    assert libro[H["titulados"]]["M9"].value == "=SUM(B9:D9)/'4. Cohortes ciclo-semestre'!B8"
    assert libro[H["cohortes"]]["AA18"].value == "=(L18-R18)/L18"
    assert [libro[H["egresados"]][c].value for c in ("B6", "K6", "A7")] == ["2022-I", "2026-II", "2017-I"]


def test_sin_tipo_de_cambio_los_bolivianos_quedan_en_blanco(datos):  # noqa: F811
    _institucional_en_bolivianos()
    libro = _libro(guia.generar(ultimo=(2026, "II")))
    assert libro[H["proyectos"]]["E9"].value == "Proyecto IDH" and libro[H["proyectos"]]["AT9"].value is None


def test_pantalla_y_descarga(client, datos):  # noqa: F811
    login(client, "inv@test.cl", PASS_TEMP)
    assert client.get("/informes/guia-antecedentes").status_code == 403
    assert client.get("/informes/guia-antecedentes/xlsx").status_code == 403
    client.post("/logout")
    login(client, "admin@test.cl", PASS_ADMIN)
    assert "Guía de Antecedentes IAC-CINDA" in client.get("/informes").text
    r = client.get("/informes/guia-antecedentes?ciclo=2026-II")
    assert r.status_code == 200 and "Qué llena el sistema" in r.text and "2022-I a 2026-II" in r.text
    r = client.get("/informes/guia-antecedentes/xlsx?ciclo=2026-II&cambio=6,96")
    assert r.headers["content-disposition"] == 'attachment; filename="guia-antecedentes-2026-II.xlsx"'
    assert _libro(r.content)[H["id"]]["B39"].value == "2026-II"
    r = client.get("/informes/guia-antecedentes/xlsx?cambio=seis")    # vuelve a la pantalla con el aviso
    assert r.headers["content-type"].startswith("text/html") and "Indique el tipo de cambio" in r.text
    assert 'formaction="/informes/guia-antecedentes/xlsx"' in client.get("/informes/guia-antecedentes").text
    assert "Indique el tipo de cambio" in client.get("/informes/guia-antecedentes?cambio=seis").text
    assert usuarios.por_email("admin@test.cl") and vinculos  # (importados para las fixtures)
