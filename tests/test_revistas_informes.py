"""Catálogo de revistas, etapa R3: revistas en los informes (ámbito, acceso abierto, cuartil, origen
del dato), modo estricto para la acreditación y sección «Datos y método»."""
import io

import pytest
from openpyxl import load_workbook

from app.comun import config, usuarios, vinculos
from app.investigacion import productos, revistas
from app.acreditacion import informes
from app.comun.db import get_db
from tests.conftest import PASS_ADMIN, PASS_TEMP, login


def _indice(nombre):
    return get_db().q("SELECT id FROM indice WHERE nombre = ?", [nombre])[0]["id"]


@pytest.fixture
def datos(app):
    emmy = usuarios.por_email("inv@test.cl")["id"]
    usuarios.cambiar_password(emmy, PASS_TEMP, temporal=False)
    vinculos.agregar(emmy, "titular", "2010-01-01")
    admin = usuarios.por_email("admin@test.cl")["id"]
    nacional = revistas.crear({"titulo": "Revista Boliviana de Matemática", "pais": "Bolivia", "acceso_abierto": 1,
                               "arbitrada": 1}, admin, verificada=True)
    revistas.agregar_indexacion(nacional, {"indice_id": _indice("SciELO"), "desde": "2015", "cuartil": "Q3",
                                           "fuente": "Portal SciELO Bolivia"}, admin)
    internacional = revistas.crear({"titulo": "Annals", "pais": "Estados Unidos", "acceso_abierto": 0}, admin,
                                   verificada=True)
    for indice, cuartil in (("Scopus", "Q2"), ("Web of Science (SCIE)", "Q1")):
        revistas.agregar_indexacion(internacional, {"indice_id": _indice(indice), "desde": "2000", "cuartil": cuartil,
                                                    "fuente": f"Lista {indice} 2026"}, admin)
    pendiente = revistas.crear({"titulo": "Revista Nueva", "pais": "Perú"}, admin)
    tipo = str(get_db().q("SELECT id FROM tipo_producto WHERE nombre = 'Artículo en revista'")[0]["id"])
    ids = {}
    for titulo, rid, revision in (("Nacional", nacional, None), ("Internacional", internacional, None),
                                  ("Declarado", pendiente, "arbitrado")):
        d, errores = productos.preparar(dict(tipo_id=tipo, titulo=titulo, anio="2026", mes="3",
                                             estado_publicacion="publicado", revision=revision or "sin_arbitraje",
                                             revista_id=str(rid)), [f"i:{emmy}"], [])
        assert not errores, errores
        ids[titulo] = productos.guardar(d, [f"i:{emmy}"], emmy, filiaciones=[1])
    return {"emmy": emmy, "ids": ids, "nacional": nacional, "internacional": internacional, "pendiente": pendiente}


def test_detalle_de_revista(datos):
    d = revistas.detalle(productos.obtener(datos["ids"]["Internacional"]))
    assert (d["ambito"], d["cuartil"], d["acceso_abierto"], d["origen"]) == (
        "Internacional", "Q1", 0, "Catálogo (revista verificada)")
    assert "Lista Scopus 2026" in d["fuentes"] and "verificado el" in d["fuentes"]
    d = revistas.detalle(productos.obtener(datos["ids"]["Nacional"]))
    assert (d["ambito"], d["acceso_abierto"]) == ("Nacional", 1)
    d = revistas.detalle(productos.obtener(datos["ids"]["Declarado"]))
    assert (d["origen"], d["fuentes"]) == ("Declarada (revista pendiente de verificar)", None)
    assert revistas.detalle({"tipo": "Libro", "revista": None})["ambito"] is None


def test_memoria_informa_las_revistas(datos):
    m = informes.memoria(2026)
    rv = m["revistas"]
    assert rv["articulos"] == 3
    assert rv["por_ambito"] == {"Nacional": 1, "Internacional": 2}  # Perú es internacional
    assert rv["por_cuartil"] == {"Q1": 1, "Q3": 1} and rv["acceso_abierto"] == 1
    assert (rv["catalogo"], rv["declarados"]) == (2, 1)
    assert rv["politica"]["arbitrado"] == ["Latindex Catálogo"] and "Scopus" in rv["politica"]["indexado"]


def test_modo_estricto_para_la_acreditacion(datos, monkeypatch):
    monkeypatch.setattr(config, "EXIGIR_REVISTA_VERIFICADA", True)
    m = informes.memoria(2026)
    assert {p["titulo"] for p in m["productos"]} == {"Nacional", "Internacional"}
    assert [(p["titulo"], p["motivo"]) for p in m["por_confirmar"]] == [
        ("Declarado", "revista sin verificar en el catálogo")]
    assert m["revistas"]["exige_verificada"]


def test_memoria_html_y_xlsx(client, datos):
    login(client, "admin@test.cl", PASS_ADMIN)
    r = client.get("/informes/memoria?anio=2026")
    assert "Revistas de los artículos" in r.text and "Datos y método" in r.text
    assert "Latindex Catálogo" in r.text and "clasificación declarada (revista sin verificar)" in r.text
    libro = load_workbook(io.BytesIO(client.get("/informes/memoria/xlsx?anio=2026").content))
    hoja = libro["Productos"]
    encabezados = [c.value for c in hoja[1]]
    for columna in ("Revista (catálogo)", "Ámbito", "Acceso abierto", "Mejor cuartil", "Origen de la clasificación",
                    "Fuente de la indexación"):
        assert columna in encabezados
    fuentes = [f[encabezados.index("Fuente de la indexación")].value for f in hoja.iter_rows(min_row=2)]
    assert any(f and "Portal SciELO Bolivia" in f for f in fuentes)
