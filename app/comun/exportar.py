"""Exportación a CSV y XLSX con las celdas neutralizadas contra la «inyección de fórmulas». La usan los
informes (app/acreditacion/informes.py) y la auditoría."""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.comun import marca


def celda_segura(valor):
    """Evita la «inyección de fórmulas»: Excel ejecuta como fórmula un texto que empieza con
    = + - @ (y openpyxl lo guardaría como fórmula). Se antepone un apóstrofo."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


def a_csv(filas: list[dict], columnas: list[str] | None = None) -> bytes:
    """CSV en UTF-8 con BOM (Excel reconoce así las tildes). Con `columnas`, lleva encabezado aunque no
    haya filas."""
    buf = io.StringIO()
    if filas or columnas:
        w = csv.DictWriter(buf, fieldnames=columnas or list(filas[0]))
        w.writeheader()
        w.writerows({k: celda_segura(v) for k, v in f.items()} for f in filas)
    return buf.getvalue().encode("utf-8-sig")


ENCABEZADO = (Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=marca.AZUL.lstrip("#").upper()))


def a_xlsx(hojas: dict[str, list[dict]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre[:31])
        if not filas:
            ws.append(["Sin datos"])
            continue
        columnas = list(filas[0])
        ws.append(columnas)
        for c in ws[1]:    # encabezado con el azul marino de la Carrera (PLAN.md §55)
            c.font, c.fill = ENCABEZADO
        for f in filas:
            ws.append([celda_segura(f[c]) for c in columnas])
        for i, c in enumerate(columnas, start=1):
            ancho = max(len(str(c)), *(len(str(f[c])) for f in filas))
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, ancho + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
